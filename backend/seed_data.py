import os
import sys
from pathlib import Path
import openpyxl
from backend.database import get_db_connection, init_db

if sys.stdout.encoding and sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass


def seed_events_if_needed():
    """Seeds the events table from data/events_seed.xlsx if the table is empty."""
    init_db()

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT COUNT(*) AS cnt FROM events")
    count = cursor.fetchone()["cnt"]

    if count > 0:
        print(f"Events table already has {count} rows. Skipping seed.")
        conn.close()
        return

    seed_file_path = Path(__file__).resolve().parent.parent / "data" / "events_seed.xlsx"

    if not seed_file_path.exists():
        conn.close()
        raise FileNotFoundError(
            f"Missing seed file at '{seed_file_path}'. "
            "Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    wb = openpyxl.load_workbook(seed_file_path)
    if "Events" not in wb.sheetnames:
        conn.close()
        raise ValueError(
            f"Sheet 'Events' not found in {seed_file_path}. Sheets found: {wb.sheetnames}"
        )

    ws = wb["Events"]

    # Read rows
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        conn.close()
        raise ValueError("Seed file contains no rows.")

    header = rows[0]
    expected_header = (
        "id", "title", "category", "description",
        "event_date", "event_time", "venue", "capacity", "fee", "icon"
    )

    # Validate header matches expected columns
    for idx, expected_col in enumerate(expected_header):
        if idx >= len(header) or str(header[idx]).strip() != expected_col:
            print(f"Warning: Header column {idx} is '{header[idx]}' vs expected '{expected_col}'")

    inserted_count = 0
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        ev_id, title, category, description, event_date, event_time, venue, capacity, fee, icon = row[:10]
        cursor.execute("""
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            int(ev_id),
            str(title).strip() if title is not None else "",
            str(category).strip() if category is not None else "",
            str(description).strip() if description is not None else None,
            str(event_date).strip() if event_date is not None else "",
            str(event_time).strip() if event_time is not None else "",
            str(venue).strip() if venue is not None else "",
            int(capacity) if capacity is not None else 0,
            int(fee) if fee is not None else 0,
            str(icon).strip() if icon is not None else ""
        ))
        inserted_count += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {inserted_count} events from {seed_file_path.name}.")


if __name__ == "__main__":
    seed_events_if_needed()

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events ORDER BY id ASC")
    rows = cursor.fetchall()
    print(f"\n--- Events in Database ({len(rows)} rows) ---")
    for r in rows:
        print(f"[{r['id']}] {r['icon']} {r['title']} | {r['category']} | {r['event_date']} {r['event_time']} | {r['venue']} | Cap: {r['capacity']} | Fee: ₹{r['fee']}")
    conn.close()
